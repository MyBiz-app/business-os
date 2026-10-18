"""The payroll run up to gross pay: the arithmetic, the stages, the locks and the export."""

import csv
import io
from datetime import date
from decimal import Decimal
from uuid import UUID, uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, text

from app.payroll.components import Flags
from app.payroll.gross import Adjustment, Component, Terms, compute_gross
from app.schedule.time_clock import Breakdown, ClockRules
from tests.conftest import AuthHeaders, add_member

FIRST, LAST = date(2026, 3, 1), date(2026, 3, 31)
NONE = Flags(False, False, False)
TAXED = Flags(True, True, True)


def _component(calc: str, amount: int | None, cap: int | None = None, **extra) -> Component:
    return Component(
        id=str(uuid4()),
        kind=extra.get("kind", "travel"),
        name="x",
        calc=calc,
        amount_minor=amount,
        cap_minor=cap,
        flags=NONE,
        valid_from=extra.get("valid_from", date(2026, 1, 1)),
        valid_to=extra.get("valid_to"),
    )


def test_hourly_pay_adds_overtime_components_and_bonus() -> None:
    hours = Breakdown(regular=10, tier1=2, tier2=1, rest_day=0, total=13)
    travel = _component("per_work_day", 2000, cap=8000)
    expired = _component("monthly", 99999, valid_to=date(2026, 2, 28))
    fine = _component("monthly", 5000, kind="fixed_deduction")
    bonus = Adjustment(id=str(uuid4()), kind="bonus", amount_minor=30000, note=None, flags=TAXED)
    result = compute_gross(
        Terms("hourly", None, 5000),
        hours,
        5,
        [travel, expired, fine],
        [bonus],
        ClockRules(),
        FIRST,
        LAST,
        {},
    )
    assert result.base_minor == 5000 * 10
    assert result.overtime_minor == 5000 * 2 * 125 // 100 + 5000 * 150 // 100  # 12500 + 7500
    assert result.components_minor == 8000  # 5 days x 20.00 = 100.00, capped at 80.00
    assert result.adjustments_minor == 30000
    assert result.deductions_minor == 5000  # off the net pay, not the gross
    assert result.gross_minor == 50000 + 20000 + 8000 + 30000
    assert result.taxable_minor == 50000 + 20000 + 30000  # travel is not taxable
    assert result.pension_base_minor == 50000 + 30000
    assert result.problems == []


def test_global_pay_follows_the_scope_and_rounds_half_up() -> None:
    result = compute_gross(
        Terms("global", 1000001, None, Decimal("50")),
        Breakdown(),
        22,
        [],
        [],
        ClockRules(),
        FIRST,
        LAST,
        {},
    )
    assert result.base_minor == 500001  # 500000.5 rounds up


def test_global_overtime_pays_only_the_overtime_hours_and_derives_the_rate() -> None:
    hours = Breakdown(regular=180, tier1=2, tier2=0, rest_day=4, total=186)
    result = compute_gross(
        Terms("global_overtime", 1820000, None), hours, 22, [], [], ClockRules(), FIRST, LAST, {}
    )
    rate = 1820000 // 182  # 100.00 an hour
    assert result.base_minor == 1820000
    assert result.overtime_minor == rate * 2 * 125 // 100 + rate * 4 * 150 // 100
    assert "derived_hourly_rate" in result.problems


def test_typed_amounts_fill_the_components_that_need_them() -> None:
    by_month = _component("by_month", None, kind="phone")
    automatic = _component("automatic", None, kind="recuperation")
    terms = Terms("global", 1000000, None)
    open_ = compute_gross(
        terms, Breakdown(), 22, [by_month, automatic], [], ClockRules(), FIRST, LAST, {}
    )
    assert f"typed_amount_missing:{by_month.id}" in open_.problems
    assert f"automatic_rate_missing:{automatic.id}" in open_.problems
    typed = compute_gross(
        terms, Breakdown(), 22, [by_month], [], ClockRules(), FIRST, LAST, {by_month.id: 7000}
    )
    assert typed.components_minor == 7000 and typed.problems == []


def open_payroll(client: TestClient, headers: dict) -> None:
    """Gives the business the payroll module and the legal details a payroll needs."""
    assert (
        client.put(
            "/tenants/current/modules", json={"modules": {"payroll": 1}}, headers=headers
        ).status_code
        == 200
    )
    details = client.patch(
        "/tenants/current",
        json={
            "legal_entity_type": "licensed_dealer",
            "business_number": "000000018",
            "address_line": "12 Herzl St",
            "address_city": "Holon",
            "employer_file_number": "940123456",
        },
        headers=headers,
    )
    assert details.status_code == 200, details.text


def _setup(client: TestClient, studio: dict, engine: Engine) -> dict:
    headers = studio["headers"]
    open_payroll(client, headers)
    coach, owner = studio["coach"], studio["owner"]
    assert (
        client.put(
            f"/staff/{coach}/employment",
            json={"employment_type": "hourly", "hourly_rate_minor": 6000},
            headers=headers,
        ).status_code
        == 200
    )
    assert (
        client.put(
            f"/staff/{owner}/employment",
            json={"employment_type": "global", "monthly_amount_minor": 1500000},
            headers=headers,
        ).status_code
        == 200
    )
    client.post(
        f"/staff/{coach}/pay-components",
        json={
            "kind": "travel",
            "name": "Travel",
            "amount_minor": 2000,
            "valid_from": "2026-01-01",
        },
        headers=headers,
    )
    # Monday 2 March 2026, 09:00-18:00 in Jerusalem: 8.6 regular, 0.4 at 125%.
    entry = client.post(
        "/time-clock/entries",
        json={
            "user_id": str(coach),
            "location_id": studio["location"]["id"],
            "clock_in": "2026-03-02T09:00:00+02:00",
            "clock_out": "2026-03-02T18:00:00+02:00",
            "reason": "test",
        },
        headers=headers,
    )
    assert entry.status_code == 201, entry.text
    return {"headers": headers, "coach": str(coach), "owner": str(owner)}


def _line(run: dict, user: str) -> dict:
    return next(line for line in run["lines"] if line["user_id"] == user)


def test_a_run_goes_through_its_stages_and_locks_the_month(
    client: TestClient, auth: AuthHeaders, studio: dict, engine: Engine
) -> None:
    ctx = _setup(client, studio, engine)
    headers, coach, owner = ctx["headers"], ctx["coach"], ctx["owner"]

    created = client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers)
    assert created.status_code == 201, created.text
    run = created.json()
    assert run["stage"] == "hours" and run["payslips_draft"] is True
    assert (
        client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers).status_code == 409
    )

    hourly = _line(run, coach)
    # 8.6 regular + 0.4 overtime at 125%, at 60.00 an hour; travel for the one day worked.
    assert hourly["base_minor"] == 51600 and hourly["overtime_minor"] == 3000
    assert hourly["components_minor"] == 2000 and hourly["gross_minor"] == 56600
    assert {w["code"] for w in hourly["warnings"]} >= {"hours_not_approved", "no_form101"}
    assert _line(run, owner)["base_minor"] == 1500000  # global: no hours needed

    # The hours are not approved yet.
    blocked = client.post(
        "/payroll/runs/2026-03/stage", json={"stage": "additions"}, headers=headers
    )
    assert blocked.status_code == 409 and blocked.json()["detail"] == "run_blocked"
    skip = client.post("/payroll/runs/2026-03/stage", json={"stage": "review"}, headers=headers)
    assert skip.status_code == 422

    assert (
        client.put(
            f"/time-clock/people/{coach}/approval", params={"month": "2026-03"}, headers=headers
        ).status_code
        == 204
    )
    moved = client.post("/payroll/runs/2026-03/stage", json={"stage": "additions"}, headers=headers)
    assert moved.status_code == 200 and moved.json()["stage"] == "additions"
    # Once the hours are locked into the run, the month's approval cannot be taken back.
    reopen = client.delete(
        f"/time-clock/people/{coach}/approval", params={"month": "2026-03"}, headers=headers
    )
    assert reopen.status_code == 409 and reopen.json()["detail"] == "payroll_locked"

    # A bonus added now is picked up when the run moves on.
    bonus = client.post(
        "/payroll/bonuses",
        json={
            "target": {"kind": "employee", "user_ids": [coach]},
            "kind": "bonus",
            "period": "2026-03-15",
            "amount_minor": 20000,
        },
        headers=headers,
    )
    assert bonus.status_code == 201, bonus.text
    review = client.post("/payroll/runs/2026-03/stage", json={"stage": "review"}, headers=headers)
    assert review.status_code == 200
    assert _line(review.json(), coach)["gross_minor"] == 56600 + 20000
    assert _line(review.json(), coach)["items"][-1]["section"] == "adjustment"

    # Warnings must be acknowledged before approval; errors cannot be.
    open_ = client.post("/payroll/runs/2026-03/stage", json={"stage": "approved"}, headers=headers)
    assert open_.status_code == 409 and open_.json()["detail"] == "warnings_open"
    for user in (coach, owner):
        for warning in _line(client.get("/payroll/runs/2026-03", headers=headers).json(), user)[
            "warnings"
        ]:
            ack = client.post(
                f"/payroll/runs/2026-03/lines/{user}/acknowledge",
                json={"key": warning["key"]},
                headers=headers,
            )
            assert ack.status_code == 200, ack.text
    approved = client.post(
        "/payroll/runs/2026-03/stage", json={"stage": "approved"}, headers=headers
    )
    assert approved.status_code == 200, approved.text
    final = approved.json()
    assert final["stage"] == "approved" and final["approved_at"] and final["payslips_draft"]
    assert final["summary"]["gross_minor"] == 1500000 + 76600
    assert final["summary"]["open_errors"] == 0

    # Locked: nothing changes the approved month.
    assert client.post("/payroll/runs/2026-03/refresh", headers=headers).status_code == 409
    late = client.post(
        "/payroll/bonuses",
        json={
            "target": {"kind": "employee", "user_ids": [coach]},
            "kind": "bonus",
            "period": "2026-03-20",
            "amount_minor": 1000,
        },
        headers=headers,
    )
    assert late.status_code == 409 and late.json()["detail"] == "period_locked"
    assert client.post("/payroll/runs/2026-03/accountant-check", headers=headers).status_code == 204
    assert client.get("/payroll/runs/2026-03", headers=headers).json()["payslips_draft"] is False

    with engine.connect() as connection:
        saved = connection.execute(
            text("SELECT export_file_id FROM app.payroll_runs WHERE period = '2026-03-01'")
        ).scalar_one()
    assert saved is not None  # the workbook is in the payroll folder

    listed = client.get("/payroll/runs", headers=headers).json()
    assert listed[0]["period"] == "2026-03-01" and listed[0]["gross_minor"] == 1576600


def test_the_export_is_a_row_per_person(client: TestClient, studio: dict, engine: Engine) -> None:
    ctx = _setup(client, studio, engine)
    headers = ctx["headers"]
    client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers)
    body = client.get("/payroll/runs/2026-03/export", params={"format": "csv"}, headers=headers)
    assert body.status_code == 200 and "payroll-2026-03.csv" in body.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(body.content.decode("utf-8-sig"))))
    assert len(rows) == 2
    hourly = next(row for row in rows if row["employment_type"] == "hourly")
    assert Decimal(hourly["gross"]) == Decimal("566.00") and Decimal(hourly["travel"]) == 20
    book = load_workbook(
        io.BytesIO(client.get("/payroll/runs/2026-03/export", headers=headers).content)
    )
    assert book.active.max_row == 3


def test_typed_amounts_the_owner_only_approval_and_who_may_look(
    client: TestClient, auth: AuthHeaders, studio: dict, engine: Engine
) -> None:
    headers, tenant_id = studio["headers"], studio["tenant_id"]
    owner = studio["owner"]
    open_payroll(client, headers)
    client.put(
        f"/staff/{owner}/employment",
        json={"employment_type": "global", "monthly_amount_minor": 1000000},
        headers=headers,
    )
    phone = client.post(
        f"/staff/{owner}/pay-components",
        json={
            "kind": "phone",
            "name": "Phone",
            "calc": "by_month",
            "valid_from": "2026-01-01",
        },
        headers=headers,
    ).json()
    run = client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers).json()
    assert f"typed_amount_missing:{phone['id']}" in [
        w["key"] for w in _line(run, str(owner))["warnings"]
    ]
    stuck = client.post("/payroll/runs/2026-03/stage", json={"stage": "additions"}, headers=headers)
    assert stuck.status_code == 200  # nothing to fix before the hours stage is left
    stuck = client.post("/payroll/runs/2026-03/stage", json={"stage": "review"}, headers=headers)
    assert stuck.status_code == 409 and stuck.json()["detail"] == "run_blocked"
    cannot = client.post(
        f"/payroll/runs/2026-03/lines/{owner}/acknowledge",
        json={"key": f"typed_amount_missing:{phone['id']}"},
        headers=headers,
    )
    assert cannot.status_code == 422
    typed = client.put(
        f"/payroll/runs/2026-03/lines/{owner}/typed",
        json={"component_id": phone["id"], "amount_minor": 4500},
        headers=headers,
    )
    assert typed.status_code == 200
    assert _line(typed.json(), str(owner))["gross_minor"] == 1004500
    assert (
        client.post(
            "/payroll/runs/2026-03/stage", json={"stage": "review"}, headers=headers
        ).status_code
        == 200
    )

    # A manager runs the months but only the owner approves; staff see nothing.
    manager, staff = uuid4(), uuid4()
    add_member(engine, tenant_id, manager, "manager")
    add_member(engine, tenant_id, staff, "staff")
    manager_headers = auth(manager, tenant_id)
    assert client.get("/payroll/runs/2026-03", headers=manager_headers).status_code == 200
    forbidden = client.post(
        "/payroll/runs/2026-03/stage", json={"stage": "approved"}, headers=manager_headers
    )
    assert forbidden.status_code == 403
    assert client.get("/payroll/runs", headers=auth(staff, tenant_id)).status_code == 403
    # The manager does not see the owner's pay.
    seen = client.get("/payroll/runs/2026-03", headers=manager_headers).json()
    assert all(line["user_id"] != str(owner) for line in seen["lines"])

    back = client.post("/payroll/runs/2026-03/stage", json={"stage": "hours"}, headers=headers)
    assert back.status_code == 200 and back.json()["stage"] == "hours"
    assert client.get("/payroll/runs/2025-01", headers=headers).status_code == 404
    assert UUID(str(tenant_id))


def test_the_server_gates_payroll_on_the_module_and_the_legal_details(
    client: TestClient, studio: dict
) -> None:
    headers = studio["headers"]
    blocked = client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers)
    assert (blocked.status_code, blocked.json()["detail"]) == (403, "module_disabled")
    client.put("/tenants/current/modules", json={"modules": {"payroll": 1}}, headers=headers)
    blocked = client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers)
    assert (blocked.status_code, blocked.json()["detail"]) == (409, "payroll_details_missing")
    assert client.get("/payroll/runs", headers=headers).status_code == 200  # looking stays open
    open_payroll(client, headers)
    assert (
        client.post("/payroll/runs", json={"period": "2026-03"}, headers=headers).status_code == 201
    )
